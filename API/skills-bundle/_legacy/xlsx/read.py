#!/usr/bin/env python3
"""
xlsx-read — Lit un fichier Excel et retourne :
- toutes les feuilles avec données
- les valeurs CALCULÉES des formules (data_only=True d'openpyxl, sinon fallback xlcalculator)
- les formules originales en parallèle
- métadonnées (auteur, dates, taille, nb feuilles)

Usage : python3 read.py <input.xlsx> [--sheet <nom>] [--max-rows N]
Output : JSON sur stdout

Stratégie pour les formules :
1. openpyxl(data_only=True) lit les valeurs cached (Excel/LibreOffice les sauve)
2. Si une cell formule n'a pas de cached value (fichier créé programmatiquement
   sans avoir été ouvert dans Excel), fallback xlcalculator pour recalculer
"""
import sys
import json
import os
from pathlib import Path


def safe_value(v):
    """Convertit valeurs Excel en JSON-safe (datetime → ISO, etc.)."""
    from datetime import datetime, date, time
    if v is None:
        return None
    if isinstance(v, (datetime, date, time)):
        return v.isoformat()
    if isinstance(v, bytes):
        try:
            return v.decode('utf-8', errors='replace')
        except Exception:
            return str(v)
    return v


def read_xlsx(path, sheet_filter=None, max_rows=10000):
    from openpyxl import load_workbook

    # 1. Workbook avec valeurs CALCULÉES (cached)
    wb_values = load_workbook(path, data_only=True, read_only=False)
    # 2. Workbook avec formules originales (pour les afficher en parallèle)
    wb_formulas = load_workbook(path, data_only=False, read_only=False)

    # Détecte les cells formules sans cached value (fichier jamais ouvert dans Excel)
    needs_recompute = False
    sample_check = 0
    for sn in wb_values.sheetnames:
        if sample_check >= 50:
            break
        ws_v = wb_values[sn]
        ws_f = wb_formulas[sn]
        for row_v, row_f in zip(ws_v.iter_rows(max_row=20), ws_f.iter_rows(max_row=20)):
            for cv, cf in zip(row_v, row_f):
                sample_check += 1
                if isinstance(cf.value, str) and cf.value.startswith('=') and cv.value is None:
                    needs_recompute = True
                    break
            if needs_recompute:
                break
        if needs_recompute:
            break

    # 3. Si des formules n'ont pas de valeur cached, recalcule via xlcalculator
    recomputed = {}
    if needs_recompute:
        try:
            from xlcalculator import ModelCompiler, Evaluator
            compiler = ModelCompiler()
            new_model = compiler.read_and_parse_archive(path)
            evaluator = Evaluator(new_model)
            for sn in wb_formulas.sheetnames:
                ws_f = wb_formulas[sn]
                for row in ws_f.iter_rows():
                    for cell in row:
                        if isinstance(cell.value, str) and cell.value.startswith('='):
                            try:
                                ref = f"{sn}!{cell.coordinate}"
                                val = evaluator.evaluate(ref)
                                recomputed[ref] = safe_value(val)
                            except Exception:
                                pass
        except Exception as e:
            recomputed['_error'] = f'xlcalculator_failed: {e}'

    sheets_out = []
    for sn in wb_values.sheetnames:
        if sheet_filter and sn != sheet_filter:
            continue
        ws_v = wb_values[sn]
        ws_f = wb_formulas[sn]
        rows = []
        formulas = []
        for r_idx, (row_v, row_f) in enumerate(zip(
                ws_v.iter_rows(max_row=max_rows),
                ws_f.iter_rows(max_row=max_rows)), start=1):
            row_values = []
            for c_idx, (cv, cf) in enumerate(zip(row_v, row_f), start=1):
                val = safe_value(cv.value)
                if val is None and isinstance(cf.value, str) and cf.value.startswith('='):
                    # Cell formule sans cached → essaye recomputed
                    ref = f"{sn}!{cf.coordinate}"
                    val = recomputed.get(ref)
                if isinstance(cf.value, str) and cf.value.startswith('='):
                    formulas.append({
                        'cell': cf.coordinate,
                        'formula': cf.value,
                        'computed': val,
                    })
                row_values.append(val)
            # Skip pure empty rows
            if any(v is not None and v != '' for v in row_values):
                rows.append(row_values)

        sheets_out.append({
            'name': sn,
            'maxRow': ws_v.max_row,
            'maxCol': ws_v.max_column,
            'rowsRead': len(rows),
            'rows': rows,
            'formulas': formulas[:200],  # limite pour éviter payload massif
        })

    file_size = os.path.getsize(path)
    props = wb_values.properties
    return {
        'ok': True,
        'fileSize': file_size,
        'sheetCount': len(wb_values.sheetnames),
        'sheetNames': wb_values.sheetnames,
        'metadata': {
            'creator': props.creator,
            'lastModifiedBy': props.lastModifiedBy,
            'created': safe_value(props.created),
            'modified': safe_value(props.modified),
            'title': props.title,
        },
        'recomputedFormulas': len(recomputed) - (1 if '_error' in recomputed else 0),
        'recomputeError': recomputed.get('_error'),
        'sheets': sheets_out,
    }


if __name__ == '__main__':
    args = sys.argv[1:]
    if not args:
        print(json.dumps({'ok': False, 'error': 'usage: read.py <file.xlsx> [--sheet NAME] [--max-rows N]'}))
        sys.exit(1)
    path = args[0]
    sheet_filter = None
    max_rows = 10000
    i = 1
    while i < len(args):
        if args[i] == '--sheet' and i + 1 < len(args):
            sheet_filter = args[i + 1]; i += 2; continue
        if args[i] == '--max-rows' and i + 1 < len(args):
            max_rows = int(args[i + 1]); i += 2; continue
        i += 1
    try:
        out = read_xlsx(path, sheet_filter, max_rows)
        print(json.dumps(out, ensure_ascii=False, default=str))
    except Exception as e:
        print(json.dumps({'ok': False, 'error': str(e)}))
        sys.exit(1)
