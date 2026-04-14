#!/usr/bin/env python3
"""Create a .xlsx workbook from a JSON spec.

Spec format (stdin):
{
  "title": "Report",
  "sheets": [
    {
      "name": "Data",
      "headers": ["Name", "Amount"],
      "rows": [["Apple", 12], ["Pear", 7]],
      "columnWidths": [20, 10],
      "styles": {
        "header": {"bold": true, "fillColor": "4472C4", "fontColor": "FFFFFF"},
        "alternateRows": {"fillColor": "F2F2F2"}
      },
      "freezeHeader": true,
      "formulas": [ {"cell": "C2", "formula": "=A2*B2"} ],
      "charts": [ {"type": "bar"|"line"|"pie", "dataRange": "A1:B5", "title": "Sales", "position": "E2"} ]
    }
  ]
}
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from _common import (  # noqa: E402
    read_spec, die, build_out_path, emit_relative, check_size,
)

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference
except ImportError as e:
    die(f"openpyxl not installed: {e}")


def _fill(color):
    if not color:
        return None
    c = str(color).lstrip("#").upper()
    if len(c) == 6:
        c = "FF" + c
    try:
        return PatternFill("solid", fgColor=c)
    except Exception:  # noqa: BLE001
        return None


def _font(spec):
    if not isinstance(spec, dict):
        return None
    kwargs = {}
    if spec.get("bold"):
        kwargs["bold"] = True
    if spec.get("italic"):
        kwargs["italic"] = True
    if spec.get("fontColor"):
        c = str(spec["fontColor"]).lstrip("#").upper()
        if len(c) == 6:
            c = "FF" + c
        kwargs["color"] = c
    if spec.get("size"):
        try:
            kwargs["size"] = float(spec["size"])
        except (TypeError, ValueError):
            pass
    return Font(**kwargs) if kwargs else None


def apply_header_style(ws, headers_row, header_style):
    if not header_style:
        return
    font = _font(header_style)
    fill = _fill(header_style.get("fillColor"))
    for col_i in range(1, headers_row + 1):
        pass
    for cell in ws[1]:
        if font:
            cell.font = font
        if fill:
            cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center")


def apply_alternate_rows(ws, start_row, end_row, col_count, fill_color):
    fill = _fill(fill_color)
    if not fill:
        return
    for r in range(start_row, end_row + 1):
        if (r - start_row) % 2 == 1:
            for c in range(1, col_count + 1):
                ws.cell(row=r, column=c).fill = fill


def add_chart(ws, chart_spec):
    ctype = (chart_spec.get("type") or "bar").lower()
    if ctype == "line":
        chart = LineChart()
    elif ctype == "pie":
        chart = PieChart()
    else:
        chart = BarChart()
    chart.title = chart_spec.get("title") or None

    data_range = chart_spec.get("dataRange")
    if not data_range:
        return
    try:
        # Reference supports string ranges via cached cell
        ref = Reference(ws, range_string=f"{ws.title}!{data_range}")
        chart.add_data(ref, titles_from_data=bool(chart_spec.get("titlesFromData", True)))
    except Exception as e:  # noqa: BLE001
        print(f"warning: chart data range failed: {e}", file=sys.stderr)
        return

    cats = chart_spec.get("categories")
    if cats:
        try:
            chart.set_categories(Reference(ws, range_string=f"{ws.title}!{cats}"))
        except Exception:  # noqa: BLE001
            pass

    pos = chart_spec.get("position") or "E2"
    ws.add_chart(chart, pos)


def build_sheet(wb, idx, sheet_spec):
    name = sheet_spec.get("name") or f"Sheet{idx+1}"
    # Excel sheet name limits: 31 chars, no []:*?/\
    name = str(name)[:31]
    for ch in "[]:*?/\\":
        name = name.replace(ch, "_")

    ws = wb.active if idx == 0 else wb.create_sheet(title=name)
    if idx == 0:
        ws.title = name

    headers = sheet_spec.get("headers") or []
    rows = sheet_spec.get("rows") or []

    if headers:
        for c_i, h in enumerate(headers, start=1):
            ws.cell(row=1, column=c_i, value=h)

    start_data_row = 2 if headers else 1
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row, start=1):
            ws.cell(row=start_data_row + r_i, column=c_i, value=val)

    # Column widths
    widths = sheet_spec.get("columnWidths") or []
    for i, w in enumerate(widths, start=1):
        try:
            ws.column_dimensions[get_column_letter(i)].width = float(w)
        except (TypeError, ValueError):
            continue

    # Styles
    styles = sheet_spec.get("styles") or {}
    if headers and styles.get("header"):
        apply_header_style(ws, len(headers), styles["header"])
    if styles.get("alternateRows") and rows:
        color = styles["alternateRows"].get("fillColor") or "F2F2F2"
        col_count = max(len(headers), max((len(r) for r in rows), default=0))
        apply_alternate_rows(ws, start_data_row, start_data_row + len(rows) - 1, col_count, color)

    # Freeze header
    if sheet_spec.get("freezeHeader") and headers:
        ws.freeze_panes = "A2"

    # Formulas
    for f in sheet_spec.get("formulas") or []:
        if not isinstance(f, dict):
            continue
        cell = f.get("cell")
        formula = f.get("formula")
        if cell and formula:
            ws[cell] = formula if str(formula).startswith("=") else f"={formula}"

    # Charts
    for c in sheet_spec.get("charts") or []:
        if isinstance(c, dict):
            try:
                add_chart(ws, c)
            except Exception as e:  # noqa: BLE001
                print(f"warning: chart failed: {e}", file=sys.stderr)


def main():
    spec = read_spec()
    if not isinstance(spec, dict):
        die("spec_must_be_object")

    title = spec.get("title") or "workbook"
    sheets = spec.get("sheets") or []
    if not isinstance(sheets, list):
        die("sheets_must_be_array")

    wb = Workbook()

    if not sheets:
        sheets = [{"name": "Sheet1", "headers": [], "rows": []}]

    for idx, sheet_spec in enumerate(sheets):
        if isinstance(sheet_spec, dict):
            build_sheet(wb, idx, sheet_spec)

    # Metadata
    try:
        wb.properties.title = str(title)
        if spec.get("author"):
            wb.properties.creator = str(spec["author"])
    except Exception:  # noqa: BLE001
        pass

    out_path = build_out_path(title, "xlsx")
    try:
        wb.save(out_path)
    except Exception as e:  # noqa: BLE001
        die(f"save_failed: {e}")

    check_size(out_path)
    emit_relative(out_path)


if __name__ == "__main__":
    try:
        main()
    except SystemExit:
        raise
    except Exception as e:  # noqa: BLE001
        die(f"unhandled_error: {e}")
