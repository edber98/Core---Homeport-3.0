#!/usr/bin/env python3
# Génère KINN_Radar.pptx — deck CLAIR, charte KINN (rose #e61982), schémas, logos (c4rbon noir).
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn
from pptx.dml.color import RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
INK = RGBColor(0x23, 0x26, 0x30)        # texte foncé (pas de fond sombre)
PINK = RGBColor(0xE6, 0x19, 0x82)
PINKL = RGBColor(0xFD, 0xE7, 0xF2)      # rose très clair (bandeaux)
PINKL2 = RGBColor(0xFB, 0xF2, 0xF7)
LIGHT = RGBColor(0xF7, 0xF8, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREY = RGBColor(0x8A, 0x93, 0xA6)
DARKTXT = RGBColor(0x3A, 0x3F, 0x4B)
BORD = RGBColor(0xEC, 0xEE, 0xF2)
HEAD = "Avenir Next"
BODY = "Helvetica Neue"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = prs.slide_width, prs.slide_height


def slide(bg=LIGHT):
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    r.fill.solid(); r.fill.fore_color.rgb = bg; r.line.fill.background(); r.shadow.inherit = False
    return s


def box(s, x, y, w, h, text, size=16, color=DARKTXT, bold=False, font=BODY,
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False, sp_after=6, line=1.05):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    for i, ln in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(sp_after); p.space_before = Pt(0)
        try: p.line_spacing = line
        except Exception: pass
        run = p.add_run(); run.text = ln
        f = run.font; f.size = Pt(size); f.bold = bold; f.italic = italic; f.name = font; f.color.rgb = color
    return tb


def rect(s, x, y, w, h, fill=None, lineclr=None, radius=True, lw=1.0):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                             Inches(x), Inches(y), Inches(w), Inches(h))
    shp.shadow.inherit = False
    if fill is None: shp.fill.background()
    else: shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if lineclr is None: shp.line.fill.background()
    else: shp.line.color.rgb = lineclr; shp.line.width = Pt(lw)
    return shp


def circle(s, x, y, d, fill=PINK):
    c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    c.shadow.inherit = False; c.fill.solid(); c.fill.fore_color.rgb = fill; c.line.fill.background()
    return c


def arrow(s, x1, y1, x2, y2, clr=PINK, w=1.6):
    cn = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    cn.line.color.rgb = clr; cn.line.width = Pt(w)
    try:
        ln = cn.line._get_or_add_ln()
        ln.append(ln.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'}))
    except Exception:
        pass
    return cn


def img(s, path, x, y, h=None, w=None):
    kw = {}
    if h: kw['height'] = Inches(h)
    if w: kw['width'] = Inches(w)
    return s.shapes.add_picture(os.path.join(HERE, path), Inches(x), Inches(y), **kw)


def brand_footer(s):
    # logo c4rbon NOIR en bas + mention
    img(s, "c4rbon-logo.png", 0.62, 6.92, h=0.42)
    box(s, 1.25, 6.92, 6, 0.42, "by C4rbon.group", size=11, color=GREY, font=BODY,
        anchor=MSO_ANCHOR.MIDDLE, sp_after=0)


# ───────────────────────── 1. TITRE (clair) ─────────────────────────
s = slide(WHITE)
rect(s, 0, 0, 13.333, 7.5, fill=WHITE, radius=False)
rect(s, 0, 0, 0.22, 7.5, fill=PINK, radius=False)
rect(s, 8.2, 0, 5.133, 7.5, fill=PINKL2, radius=False)   # panneau rose clair à droite
img(s, "logo-kinn.png", 0.9, 1.45, h=0.95)
box(s, 0.9, 2.8, 9, 1.4, "Radar", size=66, color=PINK, bold=True, font=HEAD, sp_after=0)
box(s, 0.92, 3.95, 7.0, 1.3,
    "L'intelligence qui surveille, audite\net optimise vos processus — en continu.",
    size=22, color=INK, font=BODY, line=1.2)
box(s, 0.92, 5.5, 7, 0.5, "Celonis + Palantir, réunis et automatisés pour la PME.",
    size=15, color=PINK, italic=True, font=BODY, sp_after=0)
# motif schématique léger dans le panneau rose
for i, lbl in enumerate(["Connecter", "Comprendre", "Analyser", "Agir"]):
    yy = 1.6 + i * 1.15
    rect(s, 9.0, yy, 3.4, 0.8, fill=WHITE, lineclr=PINK, lw=1.2)
    box(s, 9.0, yy, 3.4, 0.8, lbl, size=15, color=PINK, bold=True, font=HEAD, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, sp_after=0)
    if i < 3:
        arrow(s, 10.7, yy + 0.8, 10.7, yy + 1.15, clr=PINK, w=1.6)
brand_footer(s)

# ───────────────────────── 2. PROBLÈME ─────────────────────────
s = slide(LIGHT)
box(s, 0.7, 0.55, 12, 0.9, "Le vrai problème", size=38, color=INK, bold=True, font=HEAD)
box(s, 0.72, 1.5, 11.8, 0.8,
    "Votre entreprise tourne sur 10 logiciels qui ne se parlent pas. Personne n'a la vision globale.",
    size=18, color=DARKTXT, font=BODY, line=1.2)
pains = [
    ("Où passe l'argent ?", "Impayés et coûts cachés, sans alerte."),
    ("Tout traîne", "Commandes et dossiers bloqués, sans qu'on sache pourquoi."),
    ("Trop tard", "Les problèmes se découvrent après coup."),
    ("Audit = des semaines", "…et déjà périmé quand il sort."),
]
for i, (t, d) in enumerate(pains):
    x = 0.7 + i * 3.05
    rect(s, x, 2.6, 2.85, 2.45, fill=WHITE, lineclr=BORD)
    circle(s, x + 0.32, 2.92, 0.62, fill=PINKL)
    box(s, x + 0.32, 2.92, 0.62, 0.62, str(i + 1), size=22, color=PINK, bold=True, font=HEAD,
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, sp_after=0)
    box(s, x + 0.28, 3.7, 2.32, 0.6, t, size=16, color=INK, bold=True, font=HEAD)
    box(s, x + 0.28, 4.18, 2.34, 0.8, d, size=12.5, color=DARKTXT, font=BODY, line=1.15)
box(s, 0.72, 5.45, 11.8, 0.6, "« Je sens que ça coince quelque part, mais impossible de mettre le doigt dessus. »",
    size=17, color=PINK, italic=True, font=BODY)
brand_footer(s)

# ───────────────────────── 3. SOLUTION ─────────────────────────
s = slide(WHITE)
box(s, 0.7, 0.55, 12, 0.9, "La solution : KINN Radar", size=38, color=INK, bold=True, font=HEAD)
box(s, 0.72, 1.45, 11.8, 0.6, "Un cerveau qui surveille votre entreprise 24h/24. Vous connectez, le reste est automatique.",
    size=17, color=DARKTXT, font=BODY)
steps = [
    ("Connecte tout", "ERP, CRM, fichiers, e-mails, production, bases de données."),
    ("Comprend", "Relie clients, devis, commandes, factures, projets, fichiers, équipes."),
    ("Analyse en continu", "Goulots, retards, impayés, ruptures de stock, doublons, anomalies."),
    ("Alerte & recommande", "Pas « voici un problème » mais « voici quoi faire »."),
    ("Répond", "Vos questions en langage normal."),
]
for i, (t, d) in enumerate(steps):
    yy = 2.3 + i * 0.92
    circle(s, 0.75, yy, 0.6, fill=PINK)
    box(s, 0.75, yy, 0.6, 0.6, str(i + 1), size=20, color=WHITE, bold=True, font=HEAD,
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, sp_after=0)
    box(s, 1.6, yy, 3.4, 0.6, t, size=18, color=PINK, bold=True, font=HEAD, anchor=MSO_ANCHOR.MIDDLE, sp_after=0)
    box(s, 5.1, yy, 7.6, 0.6, d, size=14.5, color=DARKTXT, font=BODY, anchor=MSO_ANCHOR.MIDDLE, sp_after=0)
brand_footer(s)

# ───────────────────────── 4. ARCHITECTURE (clair) ─────────────────────────
s = slide(LIGHT)
box(s, 0.7, 0.5, 12, 0.9, "Comment ça marche", size=36, color=INK, bold=True, font=HEAD)
srcs = ["ERP / Dolibarr", "CRM", "Fichiers / Nextcloud", "E-mails", "Production", "Bases de données"]
for i, sname in enumerate(srcs):
    yy = 1.7 + i * 0.78
    rect(s, 0.7, yy, 2.9, 0.6, fill=WHITE, lineclr=PINK, lw=1.1)
    box(s, 0.85, yy, 2.7, 0.6, sname, size=13, color=DARKTXT, font=BODY, anchor=MSO_ANCHOR.MIDDLE, sp_after=0)
    # flèche vers le bord GAUCHE du bloc KINN, à des hauteurs étagées (faisceau lisible)
    arrow(s, 3.6, yy + 0.3, 5.12, 2.85 + i * 0.34, clr=RGBColor(0xF0, 0x8C, 0xC0), w=1.3)
rect(s, 5.15, 2.7, 3.0, 2.1, fill=PINK)
box(s, 5.15, 2.92, 3.0, 0.6, "KINN", size=26, color=WHITE, bold=True, font=HEAD, align=PP_ALIGN.CENTER, sp_after=0)
box(s, 5.2, 3.5, 2.9, 1.2, "Mémoire unifiée\n+ moteurs d'IA\n(graphe • process • modèles)",
    size=13, color=PINKL, font=BODY, align=PP_ALIGN.CENTER, line=1.2, sp_after=0)
outs = ["Vision globale", "Goulots & retards", "Alertes & actions", "Réponses en langage naturel"]
for i, o in enumerate(outs):
    yy = 1.95 + i * 0.95
    rect(s, 9.7, yy, 3.05, 0.7, fill=WHITE, lineclr=PINK, lw=1.3)
    box(s, 9.85, yy, 2.85, 0.7, o, size=13.5, color=INK, bold=True, font=BODY, anchor=MSO_ANCHOR.MIDDLE, sp_after=0)
    arrow(s, 8.15, 2.95 + i * 0.45, 9.68, yy + 0.35, clr=RGBColor(0xF0, 0x8C, 0xC0), w=1.3)
brand_footer(s)

# ───────────────────────── 5. 100% DYNAMIQUE ─────────────────────────
s = slide(WHITE)
box(s, 0.7, 0.55, 12.2, 0.9, "100 % dynamique — zéro paramétrage métier", size=32, color=INK, bold=True, font=HEAD)
box(s, 0.72, 1.5, 11.9, 0.6, "KINN n'impose rien. Il découvre ce que fait VOTRE entreprise.", size=18, color=DARKTXT, font=BODY)
feats = [
    ("Mapping automatique par IA", "Un logiciel inconnu (même une base de données) est compris par le LLM, qui apprend le schéma depuis vos données."),
    ("Ontologie qui s'adapte", "Il crée les types métier spécifiques : industrie, formation, expert-comptable…"),
    ("Analyses auto-activées", "Pas de stock → pas d'analyse stock. De la production → analyse production."),
    ("Predict-or-ask", "En cas de doute, il demande confirmation plutôt que deviner faux."),
]
for i, (t, d) in enumerate(feats):
    x = 0.7 + (i % 2) * 6.1
    yy = 2.4 + (i // 2) * 2.05
    rect(s, x, yy, 5.85, 1.8, fill=LIGHT, lineclr=BORD)
    rect(s, x, yy, 0.12, 1.8, fill=PINK, radius=False)
    box(s, x + 0.35, yy + 0.25, 5.3, 0.6, t, size=17, color=PINK, bold=True, font=HEAD)
    box(s, x + 0.35, yy + 0.82, 5.3, 0.9, d, size=13.5, color=DARKTXT, font=BODY, line=1.2)
brand_footer(s)

# ───────────────────────── 6. MÉMOIRE / RELATIONS ─────────────────────────
s = slide(LIGHT)
box(s, 0.7, 0.55, 12.2, 0.9, "La mémoire : un graphe vivant de votre entreprise", size=30, color=INK, bold=True, font=HEAD)
box(s, 0.72, 1.45, 11.9, 0.6, "Toutes vos entités et leurs relations, à travers tous vos logiciels.", size=17, color=DARKTXT, font=BODY)
cx, cy, cd = 6.35, 3.95, 1.15
nodes = [("Devis", 2.4, 2.25), ("Commande", 9.5, 2.25), ("Facture", 10.4, 4.0),
         ("Projet", 9.0, 5.55), ("Fichiers", 3.4, 5.55), ("E-mails", 1.9, 4.0)]
for t, x, y in nodes:
    arrow(s, cx + 0.55, cy + 0.55, x + 0.85, y + 0.35, clr=RGBColor(0xCC, 0x8F, 0xB4), w=1.6)
for t, x, y in nodes:
    rect(s, x, y, 1.7, 0.7, fill=WHITE, lineclr=PINK, lw=1.2)
    box(s, x, y, 1.7, 0.7, t, size=13, color=INK, bold=True, font=BODY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, sp_after=0)
circle(s, cx, cy, cd, fill=PINK)
box(s, cx, cy, cd, cd, "Client", size=15, color=WHITE, bold=True, font=HEAD, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, sp_after=0)
box(s, 0.72, 6.4, 11.5, 0.5, "Chaque élément est typé, daté, attribué — et filtrable (client, secteur, personne, période).",
    size=13.5, color=GREY, italic=True, font=BODY, sp_after=0)
brand_footer(s)

# ───────────────────────── 7. PROCESS MINING ─────────────────────────
s = slide(WHITE)
box(s, 0.7, 0.55, 12.2, 0.9, "Vos processus, cartographiés automatiquement", size=30, color=INK, bold=True, font=HEAD)
box(s, 0.72, 1.45, 11.9, 0.6, "KINN reconstruit vos processus réels (façon Celonis) et détecte où ça bloque.", size=17, color=DARKTXT, font=BODY)
flow = ["Devis", "Commande", "Livraison", "Facture", "Paiement"]
y = 2.5; fx = 0.8
for i, st in enumerate(flow):
    rect(s, fx, y, 2.0, 0.85, fill=(PINK if i == 0 else LIGHT), lineclr=PINK, lw=1.2)
    box(s, fx, y, 2.0, 0.85, st, size=14, color=(WHITE if i == 0 else INK), bold=True, font=HEAD,
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, sp_after=0)
    if i < len(flow) - 1:
        arrow(s, fx + 2.0, y + 0.42, fx + 2.4, y + 0.42, clr=PINK, w=2.0)
        box(s, fx + 1.75, y + 0.92, 0.9, 0.3, ["3 j", "5 j", "12 j", "8 j"][i], size=11, color=DARKTXT, bold=True, font=BODY, align=PP_ALIGN.CENTER, sp_after=0)
    fx += 2.4
rect(s, 0.8, 4.35, 11.75, 1.0, fill=PINKL)
box(s, 1.05, 4.5, 11.3, 0.8,
    "Goulot détecté : « Livraison → Facture : 12 jours en moyenne (secteur industrie). »\nRuptures de flux apprises : 8 factures sans commande en amont.",
    size=15, color=RGBColor(0x8A, 0x10, 0x4E), bold=True, font=BODY, line=1.25)
box(s, 0.72, 5.6, 11.9, 0.9,
    "Découverte dynamique de TOUS vos processus (vente, achat, SAV, production) — délai moyen à chaque étape, "
    "comparaison par secteur, standardisation et optimisation continues.",
    size=14, color=DARKTXT, font=BODY, line=1.25)
brand_footer(s)

# ───────────────────────── 8. MOTEURS D'IA (clair) ─────────────────────────
s = slide(LIGHT)
box(s, 0.7, 0.55, 12, 0.9, "Les moteurs d'IA", size=36, color=INK, bold=True, font=HEAD)
box(s, 0.72, 1.5, 11.9, 0.6, "Plusieurs modèles spécialisés qui apprennent de vos données et s'adaptent en continu.",
    size=17, color=DARKTXT, font=BODY)
engines = [
    ("LLM", "Mapping, typage, vérification d'anomalies, réponses en langage naturel."),
    ("Risque d'impayé", "Régression logistique par facture, ré-entraînée en continu."),
    ("Doublons & corrélation", "Similarité (kNN) cross-système : mêmes clients/personnes, IDs différents."),
    ("Séries temporelles", "Anomalies capteurs/production, prévisions."),
    ("Process mining", "Enchaînements, variantes, durées, goulots."),
    ("Adaptation", "Dérive de schéma → re-mapping automatique."),
]
for i, (t, d) in enumerate(engines):
    x = 0.7 + (i % 3) * 4.1
    yy = 2.5 + (i // 3) * 2.05
    rect(s, x, yy, 3.85, 1.8, fill=WHITE, lineclr=BORD)
    circle(s, x + 0.3, yy + 0.3, 0.5, fill=PINK)
    box(s, x + 0.3, yy + 0.3, 0.5, 0.5, "IA", size=12, color=WHITE, bold=True, font=HEAD, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, sp_after=0)
    box(s, x + 0.95, yy + 0.3, 2.8, 0.6, t, size=15.5, color=INK, bold=True, font=HEAD, anchor=MSO_ANCHOR.MIDDLE, sp_after=0)
    box(s, x + 0.3, yy + 0.95, 3.3, 0.8, d, size=12, color=DARKTXT, font=BODY, line=1.18)
brand_footer(s)

# ───────────────────────── 9. IA CONVERSATIONNELLE ─────────────────────────
s = slide(WHITE)
box(s, 0.7, 0.55, 12.2, 0.9, "Posez vos questions. En langage normal.", size=32, color=INK, bold=True, font=HEAD)
qa = [
    ("Quelles factures relancer en priorité ?",
     "24 factures impayées = 79 450 €, dont 8 sans commande en amont. Taux de recouvrement : 12 %."),
    ("Où sont mes goulots sur les clients industriels ?",
     "Livraison → Facture : 12 j (secteur fabrication). 15 tâches en retard de plus de 500 jours."),
]
y = 1.85
for q, a in qa:
    rect(s, 0.9, y, 11.5, 0.75, fill=PINK)
    box(s, 1.2, y, 11, 0.75, "?   " + q, size=16, color=WHITE, bold=True, font=BODY, anchor=MSO_ANCHOR.MIDDLE, sp_after=0)
    rect(s, 0.9, y + 0.85, 11.5, 1.05, fill=LIGHT, lineclr=PINK, lw=1.2)
    box(s, 1.2, y + 0.85, 11, 1.05, a, size=15, color=DARKTXT, font=BODY, anchor=MSO_ANCHOR.MIDDLE, line=1.25, sp_after=0)
    y += 2.35
box(s, 0.72, 6.55, 11.9, 0.5, "Réponses chiffrées et sourcées — l'IA s'appuie sur le graphe, le process mining et les modèles. Rien d'inventé.",
    size=13, color=GREY, italic=True, font=BODY, sp_after=0)
brand_footer(s)

# ───────────────────────── 10. AVANT / APRÈS ─────────────────────────
s = slide(LIGHT)
box(s, 0.7, 0.55, 12, 0.9, "Ce que ça change", size=38, color=INK, bold=True, font=HEAD)
rows = [
    ("Données éparpillées", "Vision globale unifiée"),
    ("Audits manuels périmés", "Audit continu, temps réel"),
    ("Goulots invisibles", "Détection + optimisation par IA"),
    ("Process implicites", "Process cartographiés & standardisés"),
    ("Risques découverts trop tard", "Points à risque auto-détectés"),
]
rect(s, 0.7, 1.7, 5.85, 0.6, fill=RGBColor(0xE7, 0xE9, 0xEE))
box(s, 0.95, 1.7, 5.4, 0.6, "AVANT", size=15, color=GREY, bold=True, font=HEAD, anchor=MSO_ANCHOR.MIDDLE, sp_after=0)
rect(s, 6.78, 1.7, 5.85, 0.6, fill=PINK)
box(s, 7.03, 1.7, 5.4, 0.6, "AVEC KINN", size=15, color=WHITE, bold=True, font=HEAD, anchor=MSO_ANCHOR.MIDDLE, sp_after=0)
for i, (a, b) in enumerate(rows):
    yy = 2.45 + i * 0.85
    rect(s, 0.7, yy, 5.85, 0.72, fill=WHITE, lineclr=BORD)
    box(s, 0.95, yy, 5.4, 0.72, a, size=14, color=GREY, font=BODY, anchor=MSO_ANCHOR.MIDDLE, sp_after=0)
    rect(s, 6.78, yy, 5.85, 0.72, fill=WHITE, lineclr=PINK, lw=1.0)
    box(s, 7.03, yy, 5.4, 0.72, b, size=14, color=INK, bold=True, font=BODY, anchor=MSO_ANCHOR.MIDDLE, sp_after=0)
brand_footer(s)

# ───────────────────────── 11. CAS D'USAGE ─────────────────────────
s = slide(WHITE)
box(s, 0.7, 0.55, 12, 0.9, "Pour qui ?", size=38, color=INK, bold=True, font=HEAD)
box(s, 0.72, 1.5, 11.9, 0.6, "PME et ETI qui jonglent avec plusieurs logiciels — partout où il y a des process, des délais et de l'argent.",
    size=17, color=DARKTXT, font=BODY)
cases = ["Industrie|production, stock, nomenclatures, OF", "Services / conseil|cycle devis → facture, charge",
         "Formation|sessions, financements, stagiaires", "Bureau d'étude|projets, CAO, livrables",
         "Négoce / matériel|ventes, stock, SAV", "Multi-logiciels|consolidation & corrélation"]
for i, c in enumerate(cases):
    x = 0.7 + (i % 3) * 4.1
    yy = 2.5 + (i // 3) * 2.0
    rect(s, x, yy, 3.85, 1.75, fill=LIGHT, lineclr=BORD)
    rect(s, x, yy, 3.85, 0.12, fill=PINK, radius=False)
    t, d = c.split("|")
    box(s, x + 0.3, yy + 0.32, 3.3, 0.6, t, size=17, color=PINK, bold=True, font=HEAD)
    box(s, x + 0.3, yy + 0.92, 3.3, 0.7, d, size=13, color=DARKTXT, font=BODY, line=1.2)
brand_footer(s)

# ───────────────────────── 12. CLOSING (clair) ─────────────────────────
s = slide(WHITE)
rect(s, 0, 0, 0.22, 7.5, fill=PINK, radius=False)
rect(s, 0, 5.4, 13.333, 2.1, fill=PINKL2, radius=False)
img(s, "logo-kinn.png", 0.9, 1.7, h=1.05)
box(s, 0.9, 3.15, 11.5, 1.2, "Connectez. Le reste est automatique.", size=42, color=INK, bold=True, font=HEAD, sp_after=0)
box(s, 0.92, 4.45, 11.5, 0.7, "Surveiller · Auditer · Détecter · Optimiser — par l'IA, en continu.",
    size=20, color=PINK, font=BODY, sp_after=0)
img(s, "c4rbon-logo.png", 0.9, 6.05, h=0.55)
box(s, 1.7, 6.05, 7, 0.55, "KINN by C4rbon.group", size=14, color=GREY, font=BODY, anchor=MSO_ANCHOR.MIDDLE, sp_after=0)

out = os.path.join(HERE, "..", "KINN_Radar.pptx")
prs.save(out)
print("OK ->", os.path.abspath(out), "| slides:", len(prs.slides._sldIdLst))
